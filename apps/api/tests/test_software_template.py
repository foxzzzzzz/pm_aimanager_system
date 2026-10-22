"""V1.4 import and live legacy-project compatibility acceptance tests."""

import contextlib
import copy
import shutil
import sqlite3
import uuid
from collections.abc import Iterator
from datetime import UTC, date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from project_manager_api.api.app import create_app
from project_manager_api.api.schemas import (
    MilestoneUpdateCreate,
    ProgressProposalCreate,
    ProjectChangeSetCreate,
)
from project_manager_api.db.base import Base
from project_manager_api.db.models import (
    ChangeProposal,
    MemberBinding,
    MilestoneRuntimeState,
    MobileUser,
    NotificationDelivery,
    ProjectVersion,
)
from project_manager_api.imports.errors import WorkbookValidationError
from project_manager_api.imports.manifest import TemplateManifest
from project_manager_api.imports.registry import ParserRegistry
from project_manager_api.services.errors import ConflictError
from project_manager_api.services.mobile import MobileService
from project_manager_api.services.notifications import NotificationService
from project_manager_api.services.projects import ProjectService
from project_manager_api.settings import AppSettings

ROOT = Path(__file__).resolve().parents[3]
OLD = ROOT / "tests/fixtures/lyra-template-v1/lyra_v1_sanitized.xlsx"
MANIFESTS = [ROOT / f"config/templates/lyra_project_spec-v{v}.yaml" for v in ("1.0", "1.1")]
# These explicit expectations are independent of production configuration.
SOFTWARE = {
    "J": ("驱动bring up", "M09", "M25", "EVT装机固件"),
    "O": ("DVT底包", "M14", "M26", "DVT底包"),
    "T": ("PVT底包", "M19", "M27", "PVT底包"),
    "AA": ("MP底包", "M26", "M28", "MP底包"),
}


@pytest.fixture
def sw_dir() -> Iterator[Path]:
    with TemporaryDirectory(dir=ROOT / "tmp") as directory:
        yield Path(directory)


@pytest.fixture
def new_workbook(sw_dir: Path) -> Path:
    """Pure insertion: retain all legacy business values and plan names."""
    path = sw_dir / "new-sanitized.xlsx"
    shutil.copyfile(OLD, path)
    workbook = load_workbook(path)
    sheet = workbook["项目进度表"]
    sheet.unmerge_cells("A1:Y1")
    for column, (name, _, _, output) in SOFTWARE.items():
        sheet.insert_cols(column_index_from_string(column))
        sheet[f"{column}2"] = name
        for row in (3, 4, 5):
            sheet[f"{column}{row}"] = "N/A"
        sheet[f"{column}11"] = "='项目团队构成'!C12"
        sheet[f"{column}12"] = "='项目团队构成'!C11"
        sheet[f"{column}13"] = "='项目团队构成'!C9&\"、\"&'项目团队构成'!C14"
        sheet[f"{column}14"] = "团队成员"
        sheet[f"{column}15"] = output
    sheet.merge_cells("A1:AC1")
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def registry() -> ParserRegistry:
    return ParserRegistry.from_manifest_paths(MANIFESTS)


@pytest.fixture
def service(sw_dir: Path, registry: ParserRegistry) -> Iterator[ProjectService]:
    engine = create_engine(f"sqlite:///{sw_dir / 'compatibility.sqlite'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield ProjectService(session, "manager", template_manifests=registry.manifests)
    engine.dispose()


def _import(
    service: ProjectService, registry: ParserRegistry, project_id: uuid.UUID, path: Path
) -> dict[str, Any]:
    return service.create_import(project_id, path.name, path.name, registry.parse(path))


def _published(service: ProjectService, registry: ParserRegistry, path: Path) -> uuid.UUID:
    identity = registry.parse(path).draft.project
    project = service.create_project(identity.code, identity.name)
    project_id = uuid.UUID(project["id"])
    imported = _import(service, registry, project_id, path)
    service.publish_import(uuid.UUID(imported["id"]), 0)
    service.session.commit()
    return project_id


@pytest.mark.parametrize("reverse", [False, True])
def test_sw01_identification_is_independent_of_registration_order(
    new_workbook: Path,
    reverse: bool,
) -> None:
    registry = ParserRegistry.from_manifest_paths(MANIFESTS[::-1] if reverse else MANIFESTS)
    assert registry.parse(OLD).draft.template_version == "1.0"
    draft = registry.parse(new_workbook).draft
    assert draft.template_version == "1.1"
    assert draft.document_version == "V1.3"
    assert [m.code for m in draft.milestones] == [f"M{i:02}" for i in range(1, 29)]
    for name, code, _, output in SOFTWARE.values():
        milestone = next(m for m in draft.milestones if m.name == name)
        assert (milestone.code, milestone.output) == (code, output)
        assert draft.active_plan.milestones[name].state == "not_applicable"
        assert milestone.assignments["R"] and milestone.assignments["A"]
        assert len(milestone.assignments["C"]) == 2


@pytest.mark.parametrize(
    "cell,value",
    [
        ("J2", None),
        ("O2", "EVT装机"),
        ("AC2", "未知节点"),
        ("AD2", "额外节点"),
    ],
)
def test_sw02_malformed_layout_is_not_parsed_as_another_version(
    new_workbook: Path,
    registry: ParserRegistry,
    cell: str,
    value: str | None,
) -> None:
    workbook = load_workbook(new_workbook)
    workbook["项目进度表"][cell] = value
    workbook.save(new_workbook)
    workbook.close()
    with pytest.raises(WorkbookValidationError, match="项目进度表"):
        registry.parse(new_workbook)


@pytest.mark.parametrize(
    "value,state",
    [
        (date(2026, 10, 22), "scheduled"),
        ("2026/10/22-2026/10/29", "scheduled"),
        (None, "tbd"),
        ("N/A", "not_applicable"),
    ],
)
def test_sw04_software_dates_and_details(
    new_workbook: Path,
    registry: ParserRegistry,
    value: Any,
    state: str,
) -> None:
    workbook = load_workbook(new_workbook)
    sheet = workbook["项目进度表"]
    for column in SOFTWARE:
        sheet[f"{column}5"] = value
        sheet[f"{column}6"] = date(2026, 10, 23)
        sheet[f"{column}7"] = 1
        sheet[f"{column}8"] = "偏差说明"
        sheet[f"{column}9"] = "风险说明"
    workbook.save(new_workbook)
    workbook.close()
    draft = registry.parse(new_workbook).draft
    for name, _, _, output in SOFTWARE.values():
        assert draft.active_plan.milestones[name].state == state
        milestone = next(m for m in draft.milestones if m.name == name)
        assert milestone.actual_completion.end_date == date(2026, 10, 23)
        assert (milestone.variance_days, milestone.variance_note, milestone.risk_note) == (
            1,
            "偏差说明",
            "风险说明",
        )
        assert milestone.output == output


@pytest.mark.parametrize("cell,value", [("J5", "2026/10/29-2026/10/22"), ("J11", "不存在的成员")])
def test_sw04_05_invalid_date_and_unknown_assignee_are_rejected(
    new_workbook: Path,
    registry: ParserRegistry,
    cell: str,
    value: str,
) -> None:
    workbook = load_workbook(new_workbook)
    workbook["项目进度表"][cell] = value
    workbook.save(new_workbook)
    workbook.close()
    with pytest.raises(ValueError):
        registry.parse(new_workbook)


def test_sw06_software_alone_can_select_latest_plan(
    new_workbook: Path,
    registry: ParserRegistry,
) -> None:
    workbook = load_workbook(new_workbook)
    sheet = workbook["项目进度表"]
    for row in sheet.iter_rows(min_row=5, max_row=5, min_col=2, max_col=29):
        for cell in row:
            cell.value = None
    sheet["J5"] = "N/A"
    workbook.save(new_workbook)
    assert registry.parse(new_workbook).draft.active_plan_name == sheet["A5"].value
    sheet["J5"] = None
    workbook.save(new_workbook)
    assert registry.parse(new_workbook).draft.active_plan_name == sheet["A4"].value
    workbook.close()


def test_sw08_10_20_upgrade_preserves_live_legacy_state_and_numbering(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    baseline = service.session.scalar(select(ProjectVersion))
    assert baseline is not None
    original_snapshot, original_hash = copy.deepcopy(baseline.snapshot), baseline.content_sha256
    completed = {"state": "scheduled", "start_date": "2026-09-01", "end_date": "2026-09-01"}
    delayed = {"state": "scheduled", "start_date": "2026-10-22", "end_date": "2026-10-22"}
    service.session.add(
        MilestoneRuntimeState(
            project_id=project_id,
            milestone_code="M09",
            actual_completion=completed,
            completion_revision=3,
            schedule=delayed,
            schedule_revision=2,
            schedule_plan_name=original_snapshot["active_plan_name"],
        )
    )
    proposal = service.create_progress_proposal(
        project_id,
        "M09",
        ProgressProposalCreate(
            base_version_number=1,
            start_date=date(2026, 10, 24),
            end_date=date(2026, 10, 24),
            reason="pending legacy proposal",
        ),
    )
    service.session.commit()
    before = service.review(project_id)
    # Loading the new registry/service itself must not modify a running old project.
    reloaded = ProjectService(service.session, "manager", template_manifests=registry.manifests)
    assert reloaded.review(project_id) == before
    imported = _import(service, registry, project_id, new_workbook)
    changes = [d for d in imported["diff"] if d["path"].startswith("milestones[")]
    assert {d["path"] for d in changes} == {f"milestones[M{i}]" for i in range(25, 29)}
    assert all(d["operation"] == "added" for d in changes)
    assert not any(d["path"].startswith("runtime_state[") for d in imported["diff"])
    service.cancel_import(uuid.UUID(imported["id"]))
    assert service.review(project_id) == before
    imported = _import(service, registry, project_id, new_workbook)
    service.publish_import(uuid.UUID(imported["id"]), 1)
    service.session.commit()
    review = service.review(project_id)
    names = {m["name"]: m for m in review["milestones"]}
    assert names["EVT装机"]["code"] == "M09"
    assert names["EVT装机"]["schedule"] == delayed
    task = next(t for t in service.dashboard(project_id)["tasks"] if t["code"] == "M09")
    assert task["actual_completion"] == completed
    for name, _, legacy_code, _ in SOFTWARE.values():
        assert names[name]["code"] == legacy_code
    assert [m["name"] for m in review["milestones"]] == [
        m.name for m in registry.parse(new_workbook).draft.milestones
    ]
    state = service.session.scalar(select(MilestoneRuntimeState))
    assert state is not None and (state.completion_revision, state.schedule_revision) == (3, 2)
    assert baseline.snapshot == original_snapshot and baseline.content_sha256 == original_hash
    again = _import(service, registry, project_id, new_workbook)
    assert again["diff"] == []
    assert service.publish_import(uuid.UUID(again["id"]), 2)["version_number"] == 2
    with pytest.raises(ConflictError):
        service.approve_proposal(uuid.UUID(proposal["id"]), 2)
    stored = service.session.get(ChangeProposal, uuid.UUID(proposal["id"]))
    assert stored is not None and stored.milestone_code == "M09" and stored.status == "pending"


@pytest.mark.parametrize("legacy", [False, True])
def test_sw12_old_template_cannot_remove_software_nodes(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
    legacy: bool,
) -> None:
    project_id = _published(service, registry, OLD if legacy else new_workbook)
    if legacy:
        imported = _import(service, registry, project_id, new_workbook)
        service.publish_import(uuid.UUID(imported["id"]), 1)
    before = service.review(project_id)
    with pytest.raises(ConflictError, match="新版|newer"):
        _import(service, registry, project_id, OLD)
    assert service.review(project_id) == before


def test_sw12_cached_legacy_import_cannot_publish_after_upgrade(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
    sw_dir: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    changed = sw_dir / "old-changed.xlsx"
    workbook = load_workbook(OLD)
    workbook["项目进度表"]["B5"] = date(2026, 10, 30)
    workbook.save(changed)
    workbook.close()
    cached = _import(service, registry, project_id, changed)
    upgrade = _import(service, registry, project_id, new_workbook)
    service.publish_import(uuid.UUID(upgrade["id"]), 1)
    with pytest.raises(ConflictError):
        service.publish_import(uuid.UUID(cached["id"]), 2)
    assert service.review(project_id)["current_version_number"] == 2


def test_sw13_custom_code_collision_is_rejected(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    milestone = registry.parse(OLD).draft.milestones[0].model_dump(mode="json")
    milestone.update(code="M25", name="自定义节点")
    snapshot = service.editable_data(project_id)
    window = {"state": "tbd", "start_date": None, "end_date": None}
    operations = [{"resource": "milestone", "op": "add", "key": "M25", "value": milestone}]
    operations.extend(
        {
            "resource": "plan",
            "op": "replace",
            "key": p["name"],
            "value": {**p, "milestones": {**p["milestones"], "自定义节点": window}},
        }
        for p in snapshot["plan_versions"]
    )
    change = service.create_change_set(
        project_id,
        ProjectChangeSetCreate(
            base_version_number=1,
            reason="custom",
            operations=operations,
        ),
    )
    service.publish_change_set(uuid.UUID(change["id"]), 1)
    before = service.review(project_id)
    with pytest.raises(ConflictError, match="编号|code"):
        _import(service, registry, project_id, new_workbook)
    assert service.review(project_id) == before


def test_sw19_same_code_in_two_projects_keeps_names_and_runtime_isolated(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    old_id = _published(service, registry, OLD)
    workbook = load_workbook(new_workbook)
    workbook["产品规格书"]["D4"] = "新项目（项目号：NEW-001）"
    workbook["项目进度表"]["J5"] = date(2026, 10, 22)
    workbook.save(new_workbook)
    workbook.close()
    new_id = _published(service, registry, new_workbook)
    service.session.add(
        MilestoneRuntimeState(
            project_id=old_id,
            milestone_code="M09",
            completion_revision=1,
            actual_completion={
                "state": "scheduled",
                "start_date": "2026-10-21",
                "end_date": "2026-10-21",
            },
        )
    )
    service.session.flush()
    old_m09 = next(m for m in service.dashboard(old_id)["tasks"] if m["code"] == "M09")
    new_m09 = next(m for m in service.dashboard(new_id)["tasks"] if m["code"] == "M09")
    assert old_m09["name"] == "EVT装机" and old_m09["actual_completion"]["end_date"]
    assert new_m09["name"] == "驱动bring up" and not new_m09["actual_completion"]["end_date"]
    settings = AppSettings(
        database_url="sqlite://",
        manifest_paths=MANIFESTS,
        import_storage_path=new_workbook.parent,
        max_import_size_bytes=100000,
    )
    notifications = NotificationService(service.session, settings, wechat=Mock(), sms=Mock())
    project = service._require_project(new_id)
    version = service._current_version(project)
    assert version is not None
    reminders = notifications._milestone_reminders(project, version.snapshot, date(2026, 10, 22))
    software = [r for r in reminders if r.object_id == "M09"]
    assert len(software) == 1 and software[0].project.id == new_id
    assert "驱动bring up" in software[0].title
    assert not any(r.object_id in {"M14", "M19", "M26"} for r in reminders)
    user = MobileUser(openid="cross-project-member", display_name="test member")
    service.session.add(user)
    service.session.flush()
    for pid in (old_id, new_id):
        target = next(t for t in service.dashboard(pid)["tasks"] if t["code"] == "M09")
        service.session.add(
            MemberBinding(
                project_id=pid,
                user_id=user.id,
                actor_id=f"mobile:{user.id}",
                member_name=target["assignments"]["R"][0],
                status="bound",
                invitation_token_hash=str(pid),
                invitation_expires_at=datetime(2030, 1, 1, tzinfo=UTC),
            )
        )
    service.session.flush()
    mobile = MobileService(service.session, settings, user)
    notifications.scan_daily(date(2026, 10, 22))
    delivered = service.session.scalar(
        select(NotificationDelivery).where(
            NotificationDelivery.project_id == new_id,
            NotificationDelivery.object_id == "M09",
            NotificationDelivery.channel == "in_app",
        )
    )
    assert delivered is not None and "驱动bring up" in delivered.payload["title"]
    assert notifications.scan_daily(date(2026, 10, 22)).created == 0
    groups = {g["project"]["id"]: g for g in mobile.list_my_tasks()}
    assert next(t for t in groups[str(old_id)]["tasks"] if t["code"] == "M09")["name"] == "EVT装机"
    assert (
        next(t for t in groups[str(new_id)]["tasks"] if t["code"] == "M09")["name"]
        == "驱动bring up"
    )
    proposal = mobile.create_milestone_proposal(
        new_id,
        "M09",
        MilestoneUpdateCreate(
            kind="completed",
            base_version_number=1,
            actual_completion_date=date(2026, 10, 22),
            reason="software completed",
        ),
    )
    assert service.list_proposals(old_id) == []
    service.approve_proposal(uuid.UUID(proposal["id"]), 1)
    new_m09 = next(t for t in service.dashboard(new_id)["tasks"] if t["code"] == "M09")
    assert new_m09["actual_completion"]["end_date"] == "2026-10-22"
    assert (
        next(t for t in service.dashboard(old_id)["tasks"] if t["code"] == "M09")[
            "actual_completion"
        ]["end_date"]
        == "2026-10-21"
    )
    notifications.scan_daily(date(2026, 10, 23))
    assert (
        service.session.scalar(
            select(NotificationDelivery).where(
                NotificationDelivery.object_id == "M09",
                NotificationDelivery.business_date == date(2026, 10, 23),
            )
        )
        is None
    )


def test_sw20_cancelled_first_import_does_not_pin_numbering(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    identity = registry.parse(OLD).draft.project
    project_id = uuid.UUID(service.create_project(identity.code, identity.name)["id"])
    old_import = _import(service, registry, project_id, OLD)
    service.cancel_import(uuid.UUID(old_import["id"]))
    imported = _import(service, registry, project_id, new_workbook)
    service.publish_import(uuid.UUID(imported["id"]), 0)
    review = service.review(project_id)
    assert [m["code"] for m in review["milestones"]] == [f"M{i:02}" for i in range(1, 29)]


def test_sw12_missing_new_manifest_cannot_allow_a_downgrade(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    project_id = _published(service, registry, new_workbook)
    service.template_manifests = registry.manifests[:1]
    with pytest.raises(ConflictError):
        _import(service, registry, project_id, OLD)


@pytest.mark.parametrize("legacy_wins", [False, True])
def test_sw20_first_publish_pins_scheme_and_stale_preview_cannot_switch_it(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
    legacy_wins: bool,
) -> None:
    identity = registry.parse(OLD).draft.project
    project_id = uuid.UUID(service.create_project(identity.code, identity.name)["id"])
    old_import = _import(service, registry, project_id, OLD)
    new_import = _import(service, registry, project_id, new_workbook)
    winner, loser = (old_import, new_import) if legacy_wins else (new_import, old_import)
    service.publish_import(uuid.UUID(winner["id"]), 0)
    baseline = service.editable_data(project_id)
    with pytest.raises(ConflictError):
        service.publish_import(uuid.UUID(loser["id"]), 1)
    assert service.editable_data(project_id) == baseline
    if legacy_wins:
        fresh = _import(service, registry, project_id, new_workbook)
        service.publish_import(uuid.UUID(fresh["id"]), 1)
        driver = next(
            m for m in service.review(project_id)["milestones"] if m["name"] == "驱动bring up"
        )
        assert driver["code"] == "M25"


def test_sw11_mapped_upgrade_previews_real_plan_changes_and_rechecks_runtime(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    original = service.editable_data(project_id)
    state = MilestoneRuntimeState(
        project_id=project_id,
        milestone_code="M09",
        schedule_revision=1,
        schedule={"state": "scheduled", "start_date": "2026-10-23", "end_date": "2026-10-23"},
        schedule_plan_name=original["active_plan_name"],
    )
    service.session.add(state)
    service.session.commit()
    workbook = load_workbook(new_workbook)
    workbook["项目进度表"]["A5"] = "新版变更计划"
    workbook["项目进度表"]["K5"] = date(2026, 10, 24)
    workbook.save(new_workbook)
    workbook.close()
    imported = _import(service, registry, project_id, new_workbook)
    assert any(d["path"] == "active_plan_name" for d in imported["diff"])
    assert any(d["path"] == "runtime_state[M09].schedule" for d in imported["diff"])
    state.schedule_revision += 1
    service.session.flush()
    with pytest.raises(ConflictError, match="runtime state changed"):
        service.publish_import(uuid.UUID(imported["id"]), 1)
    assert service.editable_data(project_id) == original
    assert state.schedule is not None and state.schedule_revision == 2


def test_sw18_restore_legacy_database_and_approve_existing_proposal(
    service: ProjectService,
    registry: ParserRegistry,
    sw_dir: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    proposal = service.create_progress_proposal(
        project_id,
        "M09",
        ProgressProposalCreate(
            base_version_number=1,
            start_date=date(2026, 10, 24),
            end_date=date(2026, 10, 24),
            reason="existing project remains operational",
        ),
    )
    service.session.commit()
    baseline = service.editable_data(project_id)
    # Restore an isolated SQLite backup without create_all or a migration.
    restored_path = sw_dir / "restored.sqlite"
    with contextlib.closing(sqlite3.connect(sw_dir / "compatibility.sqlite")) as source:
        with contextlib.closing(sqlite3.connect(restored_path)) as destination:
            source.backup(destination)
    engine = create_engine(f"sqlite:///{restored_path}")
    with Session(engine) as restored:
        updated = ProjectService(restored, "manager", template_manifests=registry.manifests)
        assert updated.editable_data(project_id) == baseline
        updated.approve_proposal(uuid.UUID(proposal["id"]), 1)
        restored.commit()
        assert updated.editable_data(project_id) == baseline
        tasks = updated.dashboard(project_id)["tasks"]
        assert len(updated.review(project_id)["milestones"]) == 24
        assert next(t for t in tasks if t["code"] == "M09")["plan"]["end_date"] == "2026-10-24"
    engine.dispose()
    # The original database was not modified by the restore exercise.
    assert (
        next(t for t in service.dashboard(project_id)["tasks"] if t["code"] == "M09")["plan"][
            "end_date"
        ]
        != "2026-10-24"
    )


def test_sw_api_upload_create_upgrade_and_downgrade_guards(
    sw_dir: Path,
    new_workbook: Path,
) -> None:
    settings = AppSettings(
        database_url=f"sqlite:///{sw_dir / 'api.sqlite'}",
        manifest_paths=MANIFESTS,
        import_storage_path=sw_dir / "imports",
        max_import_size_bytes=1000000,
        admin_api_token="software-test",
        phone_hmac_key="test-only",
    )
    app = create_app(settings)
    Base.metadata.create_all(app.state.engine)
    counter = 0

    def request(
        client: TestClient,
        url: str,
        *,
        path: Path | None = None,
        payload: dict[str, Any] | None = None,
    ) -> Any:
        nonlocal counter
        counter += 1
        headers = {"Authorization": "Bearer software-test", "X-Idempotency-Key": f"sw-{counter}"}
        if path is not None:
            with path.open("rb") as source:
                return client.post(url, headers=headers, files={"file": (path.name, source)})
        return client.post(url, headers=headers, json=payload)

    with TestClient(app) as client:
        created = request(client, "/api/v1/imports", path=OLD)
        assert created.status_code == 201, created.text
        project_id = created.json()["project"]["id"]
        first_id = created.json()["import"]["id"]
        assert (
            request(
                client,
                f"/api/v1/imports/{first_id}/publish",
                payload={"expected_project_version": 0},
            ).status_code
            == 200
        )
        upgraded = request(client, f"/api/v1/projects/{project_id}/imports", path=new_workbook)
        assert upgraded.status_code == 201, upgraded.text
        driver = next(d for d in upgraded.json()["diff"] if d["path"] == "milestones[M25]")
        assert driver["after"]["name"] == "驱动bring up"
        upgrade_id = upgraded.json()["id"]
        assert (
            request(
                client,
                f"/api/v1/imports/{upgrade_id}/publish",
                payload={"expected_project_version": 1},
            ).status_code
            == 200
        )
        rejected = request(client, f"/api/v1/projects/{project_id}/imports", path=OLD)
        assert rejected.status_code == 409, rejected.text
        workbook = load_workbook(new_workbook)
        workbook["产品规格书"]["D4"] = "新项目（项目号：API-NEW）"
        workbook.save(new_workbook)
        workbook.close()
        created = request(client, "/api/v1/imports", path=new_workbook)
        assert created.status_code == 201, created.text
        nodes = next(
            d["after"] for d in created.json()["import"]["diff"] if d["path"] == "milestones"
        )
        assert next(m for m in nodes if m["name"] == "驱动bring up")["code"] == "M09"
    app.state.engine.dispose()


def test_review_stale_upgrade_preview_cannot_be_rebased_to_overwrite_newer_plan(
    service: ProjectService,
    registry: ParserRegistry,
    new_workbook: Path,
    sw_dir: Path,
) -> None:
    project_id = _published(service, registry, OLD)
    upgrade = _import(service, registry, project_id, new_workbook)
    later = sw_dir / "later-old-template.xlsx"
    workbook = load_workbook(OLD)
    workbook["项目进度表"]["B5"] = date(2026, 10, 23)
    workbook.save(later)
    workbook.close()
    winner = _import(service, registry, project_id, later)
    service.publish_import(uuid.UUID(winner["id"]), 1)
    service.session.commit()
    before = service.editable_data(project_id)
    # A caller supplying the new version must not rebase an already-previewed draft.
    with pytest.raises(ConflictError):
        service.publish_import(uuid.UUID(upgrade["id"]), 2)
    assert service.editable_data(project_id) == before


def test_review_milestone_replace_cannot_change_stable_code(
    service: ProjectService,
    registry: ParserRegistry,
) -> None:
    project_id = _published(service, registry, OLD)
    before = service.editable_data(project_id)
    node = next(m for m in before["milestones"] if m["code"] == "M09")
    with pytest.raises(ConflictError, match="key|编号"):
        service.create_change_set(
            project_id,
            ProjectChangeSetCreate(
                base_version_number=1,
                reason="edit output, not identity",
                operations=[
                    {
                        "op": "replace",
                        "resource": "milestone",
                        "key": "M09",
                        "value": {**node, "code": "M99"},
                    },
                ],
            ),
        )
    assert service.editable_data(project_id) == before


@pytest.mark.parametrize("alias", [None, "M25"])
def test_review_legacy_code_mapping_must_be_complete_and_unique(
    registry: ParserRegistry,
    alias: str | None,
) -> None:
    data = registry.manifests[1].model_dump()
    node = next(m for m in data["milestones"] if m["name"] == "DVT底包")
    node["legacy_code"] = alias
    # M25 is already the bring-up node's legacy identity; two new nodes must
    # never publish with the same code even though the source codes are unique.
    with pytest.raises(ValueError, match="legacy"):
        TemplateManifest.model_validate(data)
